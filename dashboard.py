import streamlit as st
import pandas as pd
import openpyxl
import os

# Try importing scraper (will fail in cloud if selenium removed)
try:
    from scraper import fetch_events
    from storage import upsert_events
    LOCAL_MODE = True
except:
    LOCAL_MODE = False


# -------------------------
# PAGE CONFIG
# -------------------------
st.set_page_config(
    page_title="Event Analytics Dashboard",
    layout="wide"
)

st.title("🎉 Event Scraper & Analytics Dashboard")

st.divider()

# -------------------------
# SCRAPE SECTION (ONLY LOCAL)
# -------------------------

if LOCAL_MODE:

    st.subheader("🔄 Run Scraper (Local Only)")

    city = st.selectbox(
        "Select City",
        ["mumbai", "jaipur", "delhi", "bangalore", "pune"]
    )

    scrape_btn = st.button("Scrape Events")

    if scrape_btn:
        with st.spinner("Scraping events... Please wait..."):
            events = fetch_events(city)

            if len(events) == 0:
                st.warning("No events found.")
            else:
                upsert_events(events)
                st.success(f"✅ {len(events)} events scraped successfully!")

else:
    st.info("Scraping is disabled in cloud environment.")

st.divider()

# -------------------------
# LOAD DATA
# -------------------------

def load_data():
    try:
        wb = openpyxl.load_workbook("events.xlsx")
        ws = wb.active

        data = []
        headers = [cell.value for cell in ws[1]]

        for row in ws.iter_rows(min_row=2, values_only=True):
            data.append(row)

        return pd.DataFrame(data, columns=headers)

    except:
        return pd.DataFrame()

df = load_data()

# -------------------------
# ANALYTICS
# -------------------------

if not df.empty:

    st.subheader("📊 Dashboard Analytics")

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric("Total Events", len(df))

    with col2:
        st.metric("Active Events", len(df[df["status"] == "Active"]))

    with col3:
        st.metric("Expired Events", len(df[df["status"] == "Expired"]))

    st.divider()

    col4, col5 = st.columns(2)

    with col4:
        st.subheader("📍 Events per City")
        st.bar_chart(df["city"].value_counts())

    with col5:
        st.subheader("🏷 Events per Category")
        st.bar_chart(df["category"].value_counts())

    st.divider()

    st.subheader("📄 Events Data")
    st.dataframe(df, use_container_width=True)

else:
    st.warning("No data available.")
